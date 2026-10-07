"""Fase 6 — dashboard y export (`03 §12`, `fase-6`)."""
import io
import uuid
from datetime import date
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlmodel import Session

from app.modules.auth.models import Usuario
from app.modules.facturacion.models import EstadoFactura
from app.modules.servicios.models import EstadoFacturacion, EstadoServicio, Servicio
from tests.test_cobranza import _cliente, _factura, _pago_json
from tests.test_facturacion import API, _auth_headers

D = Decimal


def test_ingresos_por_mes_de_emision_con_invariante(client: TestClient, session: Session, admin_user: Usuario):
    h = _auth_headers(client, admin_user.email)
    c = _cliente(session)
    f1 = _factura(session, c, 1, "1650.00", 5)  # sept 2026
    _factura(session, c, 2, "1650.00", 6)
    client.post(f"{API}/pagos", json=_pago_json(c, "1000.00", aplicaciones=[{"factura_id": str(f1.id), "monto": "1000.00"}]), headers=h)

    r = client.get(f"{API}/dashboard/ingresos?anio=2026&mes=9", headers=h).json()
    assert r["periodo"] == {"mes": "2026-09", "facturado": "3696.00", "cobrado": "1000.00", "por_cobrar": "2696.00"}
    assert len(r["por_mes"]) == 12
    for fila in r["por_mes"]:
        assert D(fila["facturado"]) == D(fila["cobrado"]) + D(fila["por_cobrar"])  # invariante
    assert r["por_mes"][8]["mes"] == "2026-09" and r["por_mes"][0]["facturado"] == "0.00"  # meses sin datos -> ceros
    assert [m["mes"] for m in r["meses_previos"]] == ["2026-08", "2026-07"]
    assert r["totales_anio"]["facturado"] == "3696.00" and "actualizado_at" in r


def test_facturas_canceladas_no_cuentan_y_pendientes_si(client: TestClient, session: Session, admin_user: Usuario):
    h = _auth_headers(client, admin_user.email)
    c = _cliente(session)
    f1 = _factura(session, c, 1, "1650.00", 5)
    f2 = _factura(session, c, 2, "1650.00", 6)
    f1.estado = EstadoFactura.cancelada
    f2.estado = EstadoFactura.cancelacion_pendiente  # vigente hasta que el SAT confirme
    session.add(f1), session.add(f2)
    session.commit()
    r = client.get(f"{API}/dashboard/ingresos?anio=2026&mes=9", headers=h).json()
    assert r["periodo"]["facturado"] == "1848.00"


def test_por_cobrar_top_y_todos(client: TestClient, session: Session, admin_user: Usuario):
    h = _auth_headers(client, admin_user.email)
    for i in range(7):
        c = _cliente(session, nombre=f"Cliente {i}")
        _factura(session, c, 10 + i, f"{1000 + 100 * i}.00", 5)
    top = client.get(f"{API}/dashboard/por-cobrar?anio=todos&limit=5", headers=h).json()
    assert top["total_clientes"] == 7 and len(top["items"]) == 5
    saldos = [D(i["saldo"]) for i in top["items"]]
    assert saldos == sorted(saldos, reverse=True)
    assert len(client.get(f"{API}/dashboard/por-cobrar?anio=todos", headers=h).json()["items"]) == 7
    assert client.get(f"{API}/dashboard/por-cobrar?anio=2030", headers=h).json()["items"] == []


def test_pendientes_por_facturar_con_anticipados_y_rep(client: TestClient, session: Session, admin_user: Usuario):
    h = _auth_headers(client, admin_user.email)
    c = _cliente(session)
    f1 = _factura(session, c, 1, "1650.00", 5)

    def servicio(numero: int, **kw) -> None:
        session.add(Servicio(numero=numero, cliente_id=c.id, unidad_id=uuid.uuid4(), operador_id=uuid.uuid4(), fecha=date(2026, 9, 1), **kw))

    servicio(50, estado=EstadoServicio.terminado)                                   # por facturar
    servicio(51, timbrado_anticipado=True)                                          # anticipado, aún sin timbrar
    servicio(52, estado=EstadoServicio.terminado, estado_facturacion=EstadoFacturacion.facturado)  # ya facturado
    servicio(53, estado=EstadoServicio.cancelado)                                   # cancelado
    session.commit()
    client.post(f"{API}/pagos", json=_pago_json(c, "1848.00", aplicaciones=[{"factura_id": str(f1.id), "monto": "1848.00"}]), headers=h)

    p = client.get(f"{API}/dashboard/pendientes", headers=h).json()
    assert p["por_facturar"] == 2 and p["por_facturar_anticipados"] == 1
    assert p["rep_por_emitir"]["pagos"] == 1 and p["rep_por_emitir"]["monto"] == "1848.00"
    assert p["rep_por_emitir"]["vence"] == "2026-11-05"
    assert p["acuses_faltantes"] == 0 and p["inspecciones_con_fallas"] == 0


def test_cancelaciones_en_espera(client: TestClient, session: Session, admin_user: Usuario, monkeypatch):
    from unittest.mock import MagicMock

    from app.modules.facturacion import facturas as facturas_service

    h = _auth_headers(client, admin_user.email)
    c = _cliente(session)
    f1 = _factura(session, c, 1, "1650.00", 5)
    f1.estado = EstadoFactura.cancelacion_pendiente
    session.add(f1)
    session.commit()
    pac = MagicMock()
    pac.get_cfdi.return_value = {"Status": "pending"}
    monkeypatch.setattr(facturas_service, "FacturamaClient", lambda: pac)
    r = client.get(f"{API}/dashboard/cancelaciones-en-espera", headers=h).json()
    assert r["total"] == 1 and r["items"][0]["tipo"] == "factura" and r["items"][0]["folio"] == "ELG1"
    pac.get_cfdi.assert_called_once()  # refrescó el estatus al cargar el widget


def test_export_con_y_sin_datos(client: TestClient, session: Session, admin_user: Usuario):
    h = _auth_headers(client, admin_user.email)
    vacio = client.get(f"{API}/export/servicios.xlsx?fecha_desde=2030-01-01", headers=h)
    ws = load_workbook(io.BytesIO(vacio.content)).active
    assert vacio.status_code == 200 and ws.max_row == 1 and ws["A1"].value == "Fecha"  # solo encabezados
    c = _cliente(session)
    _factura(session, c, 1, "1650.00", 5)
    lleno = load_workbook(io.BytesIO(client.get(f"{API}/export/servicios.xlsx", headers=h).content)).active
    assert lleno.max_row == 2
    fila = {lleno.cell(1, i).value: lleno.cell(2, i).value for i in range(1, 18)}
    assert fila["Cliente"] == "Chemital" and fila["Total"] == 1848.0 and fila["Estado de cobro"] == "pendiente"


def test_dashboard_requiere_admin(client: TestClient, operador_user: Usuario):
    h = _auth_headers(client, operador_user.email)
    assert client.get(f"{API}/dashboard/pendientes", headers=h).status_code == 403
